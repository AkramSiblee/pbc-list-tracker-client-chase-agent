"""
build_master_tracker.py

Assembles the Master Tracker described in docs/io-spec.md B1 from:
  1. output/extracted_items.json (static audit program items — run
     extract_pbc_items.py first)
  2. sample_data/fieldwork/Interim_Testing_Sample_Requests.xlsx (the dynamic,
     fieldwork-generated stream)
  3. output/config/*.json (materiality thresholds, dependency map)

This produces a local Excel mirror for testing and review — the live version
of this tracker is a Google Sheet maintained via the Drive/Sheets connector
per docs/skill-workflow.md; this script exists so the tracker's shape can be
inspected and unit-tested without any live connector access.

Usage:
    python scripts/extract_pbc_items.py       # first
    python scripts/build_master_tracker.py    # then this
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from utils import canonical_entity  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
EXTRACTED_PATH = ROOT / "output" / "extracted_items.json"
SAMPLE_REQUESTS_PATH = ROOT / "sample_data" / "fieldwork" / "Interim_Testing_Sample_Requests.xlsx"
MATERIALITY_PATH = ROOT / "output" / "config" / "materiality_thresholds.json"
DEPENDENCY_PATH = ROOT / "output" / "config" / "dependency_map.json"
OUT_PATH = ROOT / "output" / "Master_Tracker.xlsx"

FONT = "Arial"
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

TRACKER_COLUMNS = [
    "item_id", "entity", "audit_area", "phase", "description", "source",
    "population_item_id", "owner", "confidentiality_tier", "data_privacy_flag",
    "materiality_tier", "depends_on", "vendor_specialist", "due_date",
    "status", "date_received", "version", "chase_count", "last_chase_date", "notes",
]


def load_static_items() -> list[dict]:
    if not EXTRACTED_PATH.exists():
        print("Run scripts/extract_pbc_items.py first.", file=sys.stderr)
        return []
    raw = json.loads(EXTRACTED_PATH.read_text(encoding="utf-8"))
    items = []
    for r in raw:
        items.append({
            "item_id": r["item_id"],
            "entity": r["entity"],
            "audit_area": r.get("audit_area"),
            "phase": None,  # not tagged at extraction time in this sample set
            "description": r.get("description") or r.get("raw_context"),
            "source": "Static Audit Program",
            "population_item_id": None,
            "owner": r.get("owner"),
            "confidentiality_tier": r.get("confidentiality"),
            "data_privacy_flag": None,
            "materiality_tier": None,
            "depends_on": None,
            "vendor_specialist": r.get("vendor_specialist"),
            "due_date": r.get("due_date"),
            "status": "Requested",
            "date_received": None,
            "version": "v1" if not r.get("needs_review") else None,
            "chase_count": 0,
            "last_chase_date": None,
            "notes": "NEEDS REVIEW — low extraction confidence" if r.get("needs_review") else None,
        })
    return items


def load_sample_items() -> list[dict]:
    if not SAMPLE_REQUESTS_PATH.exists():
        return []
    wb = openpyxl.load_workbook(SAMPLE_REQUESTS_PATH, data_only=True)
    ws = wb.active
    header_row = 4
    headers = [str(c.value).strip().lower().replace(" ", "_").replace("/", "_") for c in ws[header_row] if c.value]
    items = []
    for row in ws.iter_rows(min_row=header_row + 1, max_col=len(headers)):
        values = [c.value for c in row]
        if all(v in (None, "") for v in values):
            continue
        record = dict(zip(headers, values))
        entity = canonical_entity(record.get("entity", "")) or record.get("entity")
        items.append({
            "item_id": record.get("item_id"),
            "entity": entity,
            "audit_area": record.get("audit_area"),
            "phase": "Interim/Final (fieldwork)",
            "description": record.get("item_description"),
            "source": "Fieldwork Sample",
            "population_item_id": record.get("population_item_id"),
            "owner": None,
            "confidentiality_tier": "Sensitive" if "gdpr" in (record.get("confidentiality__privacy_notes") or "").lower() else "Standard",
            "data_privacy_flag": record.get("confidentiality__privacy_notes") if "gdpr" in (record.get("confidentiality__privacy_notes") or "").lower() else None,
            "materiality_tier": None,
            "depends_on": record.get("population_item_id"),
            "vendor_specialist": None,
            "due_date": record.get("due_date"),
            "status": "Requested",
            "date_received": None,
            "version": "v1",
            "chase_count": 0,
            "last_chase_date": None,
            "notes": f"Sampled references: {record.get('sampled_references')}",
        })
    return items


def apply_materiality(items: list[dict]) -> None:
    if not MATERIALITY_PATH.exists():
        return
    materiality_rows = json.loads(MATERIALITY_PATH.read_text(encoding="utf-8"))
    entity_notes = {}
    for row in materiality_rows:
        entity_name = row.get("entity", "")
        # Entity Roster/Materiality tab labels the parent as "... (Group)"; strip that for matching
        entity_key = canonical_entity(entity_name.split(" (Group)")[0].split(" (")[0]) or entity_name
        note = (row.get("basis___notes") or "")
        entity_notes[entity_key] = note

    for item in items:
        note = entity_notes.get(item["entity"], "")
        if "specific risk" in note.lower() or "specific-risk" in note.lower():
            item["materiality_tier"] = "Specific-Risk (Regardless of Materiality)"
        else:
            # Without live financial data to compare against, default new items
            # to Above Performance Materiality (the stricter/safer default) and
            # flag for the engagement team to confirm — never silently assume
            # something is immaterial.
            item["materiality_tier"] = "Above Performance Materiality (default — confirm)"


def apply_dependencies(items: list[dict]) -> None:
    if not DEPENDENCY_PATH.exists():
        return
    deps = json.loads(DEPENDENCY_PATH.read_text(encoding="utf-8"))
    dep_lookup: dict[tuple[str, str], list[dict]] = {}
    for d in deps:
        dep_lookup.setdefault((d["item_id"], d["entity"]), []).append(d)
    received_lookup = {(i["item_id"], i["entity"]) for i in items if i["status"] == "Received"}

    for item in items:
        key = (item["item_id"], item["entity"])
        item_deps = dep_lookup.get(key, [])
        if not item_deps:
            continue
        item["depends_on"] = ", ".join(d["depends_on_item_id"] for d in item_deps)
        unmet = [d for d in item_deps if (d["depends_on_item_id"], d["depends_on_entity"]) not in received_lookup]
        if unmet:
            item["status"] = "Not Yet Requested"
            waiting_on = "; ".join(f"{d['depends_on_item_id']} ({d['depends_on_entity']})" for d in unmet)
            item["notes"] = (item.get("notes") or "") + f" [Dependency-held: waiting on {waiting_on}]"


def write_tracker(items: list[dict]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Master Tracker"

    for i, col in enumerate(TRACKER_COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=col.replace("_", " ").title())
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center")

    for r, item in enumerate(items, start=2):
        for c_idx, col in enumerate(TRACKER_COLUMNS, start=1):
            cell = ws.cell(row=r, column=c_idx, value=item.get(col))
            cell.font = Font(name=FONT, size=10)
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for i in range(1, len(TRACKER_COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 18
    ws.freeze_panes = "A2"

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)


def main():
    items = load_static_items() + load_sample_items()
    apply_materiality(items)
    apply_dependencies(items)
    write_tracker(items)
    print(f"{len(items)} tracker row(s) written to {OUT_PATH}")
    held = [i["item_id"] for i in items if i["status"] == "Not Yet Requested"]
    if held:
        print(f"Dependency-held on init: {held}")


if __name__ == "__main__":
    main()
