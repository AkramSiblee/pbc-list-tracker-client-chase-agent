"""
extract_pbc_items.py

Parses every audit program in sample_data/audit_programs/ (Word, Excel, PDF —
format varies by entity, as in a real group audit) and produces a normalized
list of PBC items with a per-item extraction confidence score.

This script only extracts what's structurally findable (Item IDs, descriptions,
owners, confidentiality tags where present). It deliberately does NOT try to
resolve every ambiguity in code — e.g. detecting that a submitted file's
*content* belongs to the wrong entity, or judging whether a scanned document's
low-confidence OCR text is trustworthy, is agent-level judgment (see
docs/skill-workflow.md Step 8, Quality-check receipt) and is out of scope here.
Anything below CONFIDENCE_THRESHOLD is still emitted, but flagged for the
Changelog rather than silently included.

Usage:
    python scripts/extract_pbc_items.py [--out output/extracted_items.json]
"""
from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
from pathlib import Path

import openpyxl
import pdfplumber

sys.path.insert(0, str(Path(__file__).resolve().parent))
from utils import find_item_ids, canonical_entity  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
AUDIT_DIR = ROOT / "sample_data" / "audit_programs"
CONFIDENCE_THRESHOLD = 0.75

RESTRICTED_MARKER = re.compile(r"restricted", re.IGNORECASE)


def extract_from_docx(path: Path, entity: str) -> list[dict]:
    """Uses pandoc to get the document as GFM markdown (forces pipe-table output,
    which is far more reliably parseable than pandoc's default grid tables), then:
    1. Pulls narrative-embedded items (sentences with an [Item ID: ...] tag)
    2. Pulls the Schedule A summary table as the authoritative structured record
       — narrative gives context, the table gives the fields.

    Note: docx tables built without a designated header-row style (as ours are)
    render in GFM as an empty header + separator, with the real header labels
    appearing as the first body row — this is a pandoc/GFM quirk, not a
    malformed document, and is handled explicitly below.
    """
    result = subprocess.run(
        ["pandoc", "-t", "gfm", str(path)],
        capture_output=True, text=True, check=True,
    )
    md = result.stdout
    items = {}

    def clean(cell: str) -> str:
        return cell.replace("**", "").replace("\\$", "$").strip()

    # Narrative pass: capture the sentence context around each Item ID mention
    for line in md.splitlines():
        ids_in_line = find_item_ids(line)
        for item_id in ids_in_line:
            if item_id not in items:
                items[item_id] = {
                    "item_id": item_id,
                    "entity": entity,
                    "description": None,
                    "owner": None,
                    "confidentiality": "Restricted" if RESTRICTED_MARKER.search(line) else None,
                    "audit_area": None,
                    "source_format": "docx",
                    "confidence": 0.6,  # narrative-only mention, no structured fields yet
                    "raw_context": clean(line)[:300],
                }

    # Table pass (Schedule A): GFM pipe table. The real header (containing
    # "Item ID") may be the table's first body row rather than its declared
    # header row — detect it by content, then read every row after it as data.
    table_rows = [ln for ln in md.splitlines() if ln.strip().startswith("|")]
    if table_rows:
        header_idx = next((i for i, ln in enumerate(table_rows) if "Item ID" in ln), None)
        if header_idx is not None:
            headers = [clean(h).lower().replace(" ", "_") for h in table_rows[header_idx].strip("|").split("|")]
            for row_line in table_rows[header_idx + 1:]:
                if set(row_line.strip()) <= set("|-: "):
                    continue  # a stray separator row
                cells = [clean(c) for c in row_line.strip("|").split("|")]
                if len(cells) != len(headers):
                    continue
                row = dict(zip(headers, cells))
                item_id = row.get("item_id")
                if not item_id:
                    continue
                existing = items.get(item_id, {"item_id": item_id, "entity": entity, "source_format": "docx", "raw_context": None})
                existing.update({
                    "description": row.get("description"),
                    "owner": row.get("suggested_owner"),
                    "audit_area": row.get("isa_area") or row.get("audit_area"),
                    "confidentiality": row.get("confidentiality") or existing.get("confidentiality"),
                    "confidence": 0.95,  # came from the structured schedule table
                })
                items[item_id] = existing

    return list(items.values())


def extract_from_xlsx(path: Path, entity: str) -> list[dict]:
    """Reads a structured checklist workbook. Header row is auto-detected the
    same way generate_config.py does it (first row that looks like short
    string headers)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    header_row = None
    for row_idx in range(1, 8):
        values = [c.value for c in ws[row_idx]]
        non_empty = [v for v in values if v not in (None, "")]
        if len(non_empty) >= 2 and all(isinstance(v, str) and len(v) < 60 for v in non_empty):
            header_row = row_idx
            break
    if header_row is None:
        return []

    headers = [str(c.value).strip().lower().replace(" ", "_") for c in ws[header_row] if c.value]
    items = []
    for row in ws.iter_rows(min_row=header_row + 1, max_col=len(headers)):
        values = [c.value for c in row]
        if all(v in (None, "") for v in values):
            continue
        record = dict(zip(headers, values))
        item_id = record.get("item_id")
        if not item_id:
            continue
        due_date = record.get("due_date")
        confidence = 0.95
        missing_fields = []
        if not due_date:
            missing_fields.append("due_date")
            confidence -= 0.25
        items.append({
            "item_id": item_id,
            "entity": entity,
            "description": record.get("pbc_item_description"),
            "owner": record.get("suggested_owner"),
            "audit_area": record.get("audit_area"),
            "confidentiality": record.get("confidentiality"),
            "due_date": due_date,
            "vendor_specialist": record.get("vendor_specialist") or None,
            "source_format": "xlsx",
            "confidence": round(max(confidence, 0.0), 2),
            "missing_fields": missing_fields,
            "raw_context": None,
        })
    return items


def extract_from_pdf(path: Path, entity: str) -> list[dict]:
    """Extracts narrative Item ID mentions and the Schedule A table from a
    text-based PDF. (A scanned/image PDF would need an OCR fallback — see
    docs/io-spec.md A1 — not exercised by this sample set.)"""
    items = {}
    with pdfplumber.open(path) as pdf:
        full_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        tables = []
        for page in pdf.pages:
            tables.extend(page.extract_tables() or [])

    for line in full_text.splitlines():
        for item_id in find_item_ids(line):
            if item_id not in items:
                items[item_id] = {
                    "item_id": item_id,
                    "entity": entity,
                    "description": None,
                    "owner": None,
                    "confidentiality": "Restricted" if RESTRICTED_MARKER.search(line) else None,
                    "audit_area": None,
                    "source_format": "pdf",
                    "confidence": 0.6,
                    "raw_context": line.strip()[:300],
                }

    for table in tables:
        if not table or not table[0]:
            continue
        header = [str(h).strip().lower().replace(" ", "_") if h else "" for h in table[0]]
        if "item_id" not in header:
            continue
        for row in table[1:]:
            cleaned_row = [(cell.replace("\n", " ").strip() if isinstance(cell, str) else cell) for cell in row]
            record = dict(zip(header, cleaned_row))
            item_id = record.get("item_id")
            if not item_id:
                continue
            item_id = re.sub(r"\s+", "", item_id)  # Item IDs never contain spaces — collapse wrap artifacts
            existing = items.get(item_id, {"item_id": item_id, "entity": entity, "source_format": "pdf", "raw_context": None})
            existing.update({
                "description": record.get("description"),
                "owner": record.get("suggested_owner"),
                "confidentiality": record.get("confidentiality") or existing.get("confidentiality"),
                "notes": record.get("notes"),
                "confidence": 0.95,
            })
            items[item_id] = existing

    return list(items.values())


ENTITY_BY_FILENAME = {
    "Parent_Co_Audit_Program.docx": "Meridian Global Holdings Inc.",
    "Meridian_Europe_GmbH_Audit_Program.xlsx": "Meridian Europe GmbH",
    "Meridian_APAC_Audit_Program.pdf": "Meridian Asia Pacific Pte Ltd",
    "Meridian_Manufacturing_Co_Audit_Program.docx": "Meridian Manufacturing Co.",
}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default=str(ROOT / "output" / "extracted_items.json"))
    args = parser.parse_args()

    all_items = []
    for filename, entity in ENTITY_BY_FILENAME.items():
        path = AUDIT_DIR / filename
        if not path.exists():
            print(f"WARNING: {filename} not found, skipping", file=sys.stderr)
            continue
        if path.suffix == ".docx":
            extracted = extract_from_docx(path, entity)
        elif path.suffix == ".xlsx":
            extracted = extract_from_xlsx(path, entity)
        elif path.suffix == ".pdf":
            extracted = extract_from_pdf(path, entity)
        else:
            continue
        for item in extracted:
            item["needs_review"] = item.get("confidence", 1.0) < CONFIDENCE_THRESHOLD or bool(item.get("missing_fields"))
        all_items.extend(extracted)
        print(f"{filename}: {len(extracted)} item(s) extracted")

    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(all_items, indent=2, default=str, ensure_ascii=False))

    flagged = [i for i in all_items if i["needs_review"]]
    print(f"\nTotal: {len(all_items)} item(s), {len(flagged)} flagged for Changelog review")
    print(f"Written to {out_path}")


if __name__ == "__main__":
    main()
